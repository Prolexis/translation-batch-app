"""
================================================================================
SISTEMA DE TRADUCCIÓN POR LOTES CON AGENTES LANGCHAIN + GEMINI (TODO EN UNO)
================================================================================
Este archivo contiene la aplicación completa:
  1. Configuración y variables de entorno.
  2. Manejo de archivos (lectura y exportación en .txt, .docx, .pdf).
  3. Control de tasa de llamadas y reintentos (RateLimiter + Exponential Backoff).
  4. Agentes de procesamiento:
     - Agente Extractor (parseo de texto por párrafos/segmentos).
     - Agente Traductor (LCEL + ChatGoogleGenerativeAI).
     - Agente Validador (heurísticas de calidad + reintentos automáticos).
     - Agente Alineador (alineación por posición o semántica con Embeddings).
  5. Orquestador del Pipeline con memoria de conversación (ConversationBufferMemory).
  6. Interfaz Web interactiva en Streamlit con verificación cruzada por colores.
================================================================================
"""

import os
import io
import re
import html
import time
import random
import logging
import zipfile
import threading
import functools
import traceback
from dataclasses import dataclass, field
from typing import List, Optional, Callable

import numpy as np
import docx
from docx import Document as DocxDocument
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit

from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser
from langchain_google_genai import ChatGoogleGenerativeAI
import streamlit as st

# Carga opcional de .env
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# Logging básico
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("translation_app")


# ==============================================================================
# 1. CONFIGURACIÓN CENTRALIZADA
# ==============================================================================
@dataclass(frozen=True)
class Settings:
    # Credenciales y modelos de Gemini
    GEMINI_API_KEY: str = os.getenv("GEMINI_API_KEY", "")
    GEMINI_MODEL: str = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite")
    GEMINI_EMBEDDING_MODEL: str = os.getenv("GEMINI_EMBEDDING_MODEL", "models/text-embedding-004")

    # Idiomas por defecto
    DEFAULT_SOURCE_LANG: str = os.getenv("DEFAULT_SOURCE_LANG", "auto")
    DEFAULT_TARGET_LANG: str = os.getenv("DEFAULT_TARGET_LANG", "en")

    # Parámetros del Agente Validador
    LENGTH_DIFF_THRESHOLD: float = float(os.getenv("LENGTH_DIFF_THRESHOLD", "0.40"))  # ±40%
    MAX_RETRIES: int = int(os.getenv("MAX_RETRIES", "2"))
    MIN_UNTRANSLATED_WORD_LEN: int = int(os.getenv("MIN_UNTRANSLATED_WORD_LEN", "4"))

    # Parámetros de Alineación
    ALIGNMENT_MODE: str = os.getenv("ALIGNMENT_MODE", "position")  # "position" | "semantic"
    SEMANTIC_SIMILARITY_THRESHOLD: float = float(os.getenv("SEMANTIC_SIMILARITY_THRESHOLD", "0.55"))

    # Rate Limiting & Backoff ante 429
    MAX_CALLS_PER_MINUTE: int = int(os.getenv("MAX_CALLS_PER_MINUTE", "30"))
    BACKOFF_BASE_SECONDS: float = float(os.getenv("BACKOFF_BASE_SECONDS", "2.0"))
    BACKOFF_MAX_RETRIES: int = int(os.getenv("BACKOFF_MAX_RETRIES", "5"))
    BACKOFF_MAX_SECONDS: float = float(os.getenv("BACKOFF_MAX_SECONDS", "60.0"))

    # Paleta de colores cíclica para verificación cruzada
    HIGHLIGHT_COLORS: tuple = field(default_factory=lambda: (
        "#FFD54F",  # amarillo
        "#81C784",  # verde
        "#64B5F6",  # azul
        "#F06292",  # rosa
        "#BA68C8",  # violeta
        "#FF8A65",  # naranja
        "#4DB6AC",  # turquesa
        "#A1887F",  # marrón
    ))

    # Extensiones de archivo soportadas
    SUPPORTED_EXTENSIONS: tuple = (".txt", ".docx", ".pdf")


settings = Settings()


# ==============================================================================
# 2. ESTRUCTURA DE DATOS Y EXCEPCIONES
# ==============================================================================
class UnsupportedFormatError(Exception):
    """Lanzada cuando la extensión del archivo no es soportada."""
    pass


class CorruptFileError(Exception):
    """Lanzada cuando el archivo no se puede parsear o está vacío/dañado."""
    pass


@dataclass
class Segment:
    """Unidad mínima de texto (párrafo) que fluye por el pipeline de agentes."""
    id: int
    original: str
    translated: str = ""
    status: str = "pendiente"          # pendiente | ok | sospechoso | error
    validation_notes: List[str] = field(default_factory=list)
    retries: int = 0
    color: str = "#FFD54F"


# ==============================================================================
# 3. RATE LIMITER & BACKOFF EXPONENCIAL
# ==============================================================================
class RateLimiter:
    """Control de llamadas salientes por minuto (thread-safe)."""
    def __init__(self, max_calls_per_minute: int = None):
        self.max_calls = max_calls_per_minute or settings.MAX_CALLS_PER_MINUTE
        self._lock = threading.Lock()
        self._timestamps: list[float] = []

    def acquire(self):
        with self._lock:
            now = time.time()
            self._timestamps = [t for t in self._timestamps if now - t < 60]
            if len(self._timestamps) >= self.max_calls:
                sleep_for = 60 - (now - self._timestamps[0]) + 0.05
                logger.info("Rate limit alcanzado, pausando %.2fs", sleep_for)
                time.sleep(max(sleep_for, 0))
                now = time.time()
                self._timestamps = [t for t in self._timestamps if now - t < 60]
            self._timestamps.append(now)


global_rate_limiter = RateLimiter()


def _is_retryable_error(exc: Exception) -> bool:
    """Determina si un error devuelto por la API o red amerita reintento."""
    msg = str(exc).lower()
    keywords = [
        "429", "resource_exhausted", "rate limit", "quota",
        "deadline exceeded", "timeout", "unavailable", "503",
    ]
    return any(k in msg for k in keywords)


def with_backoff(max_retries: int = None, base_seconds: float = None, max_seconds: float = None):
    """Decorador con backoff exponencial y jitter para llamadas a la API."""
    max_retries = max_retries if max_retries is not None else settings.BACKOFF_MAX_RETRIES
    base_seconds = base_seconds if base_seconds is not None else settings.BACKOFF_BASE_SECONDS
    max_seconds = max_seconds if max_seconds is not None else settings.BACKOFF_MAX_SECONDS

    def decorator(func):
        @functools.wraps(func)
        def wrapper(*args, **kwargs):
            attempt = 0
            while True:
                try:
                    global_rate_limiter.acquire()
                    return func(*args, **kwargs)
                except Exception as exc:  # noqa: BLE001
                    attempt += 1
                    if attempt > max_retries or not _is_retryable_error(exc):
                        logger.error("Fallo definitivo tras %d intentos: %s", attempt, exc)
                        raise
                    delay = min(base_seconds * (2 ** (attempt - 1)), max_seconds)
                    delay += random.uniform(0, delay * 0.25)
                    logger.warning(
                        "Error transitorio (%s). Reintento %d/%d en %.1fs",
                        exc, attempt, max_retries, delay,
                    )
                    time.sleep(delay)
        return wrapper
    return decorator


# ==============================================================================
# 4. EXTRACCIÓN Y EXPORTACIÓN DE ARCHIVOS (.TXT, .DOCX, .PDF)
# ==============================================================================
def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_txt(file_bytes: bytes) -> List[str]:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="ignore")
    text = _clean_text(text)
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs and text:
        paragraphs = [text]
    return paragraphs


def extract_docx(file_bytes: bytes) -> List[str]:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise CorruptFileError(f"No se pudo leer el archivo .docx: {exc}") from exc
    paragraphs = []
    for p in doc.paragraphs:
        cleaned = _clean_text(p.text)
        if cleaned:
            paragraphs.append(cleaned)
    if not paragraphs:
        raise CorruptFileError("El documento .docx no contiene texto extraíble.")
    return paragraphs


def extract_pdf(file_bytes: bytes) -> List[str]:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as exc:
        raise CorruptFileError(f"No se pudo leer el archivo .pdf: {exc}") from exc
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise CorruptFileError("El PDF está protegido/cifrado con contraseña.") from exc

    paragraphs = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        page_text = _clean_text(page_text)
        for block in page_text.split("\n\n"):
            block = block.strip()
            if block:
                paragraphs.append(block)
    if not paragraphs:
        raise CorruptFileError("No se pudo extraer texto del PDF (posible PDF escaneado sin OCR).")
    return paragraphs


def _group_paragraphs(raw_paragraphs: List[str], max_chars: int = 1200) -> List[str]:
    """Agrupa párrafos pequeños para optimizar la velocidad y llamadas a la API."""
    if not raw_paragraphs:
        return []
    grouped = []
    current = []
    current_len = 0

    for p in raw_paragraphs:
        p_len = len(p)
        if current and (current_len + p_len + 2 > max_chars):
            grouped.append("\n\n".join(current))
            current = [p]
            current_len = p_len
        else:
            current.append(p)
            current_len += p_len + 2

    if current:
        grouped.append("\n\n".join(current))

    return grouped


def extract_text_segments(filename: str, file_bytes: bytes) -> List[str]:
    ext = "." + filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == ".txt":
        raw = extract_txt(file_bytes)
    elif ext == ".docx":
        raw = extract_docx(file_bytes)
    elif ext == ".pdf":
        raw = extract_pdf(file_bytes)
    else:
        raise UnsupportedFormatError(f"Formato no soportado: {ext or 'desconocido'}")
    return _group_paragraphs(raw, max_chars=1200)


def export_txt(segments: List[Segment]) -> bytes:
    text = "\n\n".join(s.translated or s.original for s in segments)
    return text.encode("utf-8")


def export_docx(segments: List[Segment]) -> bytes:
    doc = DocxDocument()
    for s in segments:
        doc.add_paragraph(s.translated or s.original)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(segments: List[Segment]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    margin = 2 * cm
    max_width = width - 2 * margin
    y = height - margin
    font_name, font_size, leading = "Helvetica", 11, 14

    c.setFont(font_name, font_size)
    for s in segments:
        paragraph = s.translated or s.original
        lines = simpleSplit(paragraph, font_name, font_size, max_width)
        for line in lines:
            if y < margin:
                c.showPage()
                c.setFont(font_name, font_size)
                y = height - margin
            c.drawString(margin, y, line)
            y -= leading
        y -= leading
    c.save()
    return buf.getvalue()


def export_segments(segments: List[Segment], target_format: str) -> bytes:
    target_format = target_format.lower().lstrip(".")
    if target_format == "txt":
        return export_txt(segments)
    if target_format == "docx":
        return export_docx(segments)
    if target_format == "pdf":
        return export_pdf(segments)
    raise UnsupportedFormatError(f"Formato de exportación no soportado: {target_format}")


# ==============================================================================
# 5. UTILIDADES DE COLOR Y VERIFICACIÓN CRUZADA
# ==============================================================================
def color_for_index(index: int) -> str:
    palette = settings.HIGHLIGHT_COLORS
    return palette[index % len(palette)]


def _contrast_text_color(hex_color: str) -> str:
    hex_color = hex_color.lstrip("#")
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return "#000000" if luminance > 0.6 else "#FFFFFF"


def render_highlighted_html(segments_text: List[str], colors: List[str]) -> str:
    blocks = []
    for i, (text, color) in enumerate(zip(segments_text, colors)):
        safe_text = html.escape(text).replace("\n", "<br>")
        fg = _contrast_text_color(color)
        tag_bg = "rgba(0,0,0,0.18)" if fg == "#000000" else "rgba(255,255,255,0.22)"
        border_col = "rgba(0,0,0,0.08)" if fg == "#000000" else "rgba(255,255,255,0.12)"
        blocks.append(
            f'<div class="segment-pill" style="background-color:{color}; color:{fg}; '
            f'padding:12px 14px; border-radius:10px; margin-bottom:10px; '
            f'font-size:0.94rem; line-height:1.55; position:relative; '
            f'border: 1px solid {border_col}; box-shadow: 0 4px 12px rgba(0,0,0,0.15); '
            f'transition: transform 0.15s ease, box-shadow 0.15s ease;">'
            f'<div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:6px;">'
            f'<span style="font-size:0.75rem; font-weight:700; background:{tag_bg}; '
            f'padding:2px 8px; border-radius:6px; letter-spacing:0.5px;">Segmento {i + 1}</span>'
            f'<span style="font-size:0.72rem; opacity:0.85;">{len(text)} caracteres</span>'
            f'</div>'
            f'<div>{safe_text}</div>'
            f'</div>'
        )
    return "\n".join(blocks)


# ==============================================================================
# 6. PIPELINE DE AGENTES
# ==============================================================================

# --- AGENTE 1: EXTRACTOR ---
class ExtractorAgent:
    name = "extractor"

    def run(self, context: dict) -> dict:
        filename = context["filename"]
        file_bytes = context["file_bytes"]

        logger.info("[Extractor] Procesando %s", filename)
        try:
            raw_paragraphs = extract_text_segments(filename, file_bytes)
        except (UnsupportedFormatError, CorruptFileError) as exc:
            logger.error("[Extractor] Error en %s: %s", filename, exc)
            context["error"] = str(exc)
            context["segments"] = []
            return context

        segments: List[Segment] = [
            Segment(id=i, original=paragraph)
            for i, paragraph in enumerate(raw_paragraphs)
        ]

        context["segments"] = segments
        context.setdefault("memory_log", []).append(
            f"[extractor] {len(segments)} segmentos extraídos de '{filename}'."
        )
        logger.info("[Extractor] %d segmentos extraídos de %s", len(segments), filename)
        return context


# --- AGENTE 2: TRADUCTOR ---
LANGUAGE_NAMES = {
    "auto": "Detección automática (auto-detect original language)",
    "es": "Español (Spanish)",
    "en": "Inglés (English)",
    "fr": "Francés (French)",
    "de": "Alemán (German)",
    "it": "Italiano (Italian)",
    "pt": "Portugués (Portuguese)",
    "zh": "Chino (Chinese)",
    "ja": "Japonés (Japanese)",
    "ru": "Ruso (Russian)",
    "ar": "Árabe (Arabic)",
}

_SYSTEM_PROMPT = (
    "Eres un traductor profesional experto. Tu única tarea es traducir el texto recibido al idioma objetivo: {target_name}.\n"
    "El idioma de origen es: {source_name}.\n"
    "Instrucciones estrictas:\n"
    "1. Traduce fielmente todo el contenido al idioma {target_name}.\n"
    "2. Devuelve ÚNICAMENTE el texto traducido, sin explicaciones, sin notas del traductor, sin introducciones y sin comillas adicionales.\n"
    "3. Conserva exactamente la estructura de párrafos y saltos de línea del texto original."
)

_RETRY_SUFFIX = (
    "\n\nIMPORTANTE: un intento previo de traducción fue rechazado por un "
    "verificador automático por el siguiente motivo: {retry_reason}. "
    "Corrige ese problema específico en esta nueva traducción."
)


class TranslatorAgent:
    name = "translator"

    def __init__(self, model_name: str = None, api_key: str = None, temperature: float = 0.2):
        self.model_name = model_name or settings.GEMINI_MODEL
        self.api_key = api_key or settings.GEMINI_API_KEY
        self._llm = ChatGoogleGenerativeAI(
            model=self.model_name,
            google_api_key=self.api_key,
            temperature=temperature,
        )
        self._parser = StrOutputParser()

    def _build_chain(self, retry_reason: str = None):
        system = (
            "You are an expert professional translator. You must translate any given source text accurately and directly into {target_name}.\n"
            "Output ONLY the translated text without any explanation, markdown commentary, or introductory phrases. "
            "Preserve original line breaks, numbers, currency symbols, proper names, and formatting."
        )
        if retry_reason:
            system += f"\n\nCRITICAL FIX: A previous translation attempt was rejected due to: {retry_reason}. Fix this specifically and ensure the output is strictly in {{target_name}}."
        prompt = ChatPromptTemplate.from_messages([
            ("system", system),
            ("human", "Translate the following source text from {source_name} into {target_name}. Return ONLY the direct translation in {target_name}:\n\n{text}"),
        ])
        return prompt | self._llm | self._parser

    @with_backoff()
    def _translate_one(self, text: str, source_lang: str, target_lang: str, retry_reason: str = None) -> str:
        source_name = LANGUAGE_NAMES.get(source_lang.lower().strip(), source_lang)
        target_name = LANGUAGE_NAMES.get(target_lang.lower().strip(), target_lang)
        chain = self._build_chain(retry_reason=retry_reason)
        result = chain.invoke({
            "text": text,
            "source_name": source_name,
            "target_name": target_name,
        })
        return str(result).strip()

    def run(self, context: dict, on_progress: Optional[Callable[[str, float], None]] = None) -> dict:
        segments: List[Segment] = context.get("segments", [])
        source_lang = context.get("source_lang", settings.DEFAULT_SOURCE_LANG)
        target_lang = context.get("target_lang", settings.DEFAULT_TARGET_LANG)
        total_segs = len(segments) or 1

        for i, seg in enumerate(segments):
            if on_progress:
                frac = 0.25 + (0.45 * (i / total_segs))
                on_progress(f"Traduciendo segmento {i + 1}/{total_segs}", frac)
            if seg.status == "ok":
                continue
            try:
                seg.translated = self._translate_one(seg.original, source_lang, target_lang)
                seg.status = "traducido"
            except Exception as exc:  # noqa: BLE001
                logger.error("[Translator] Fallo traduciendo segmento %d: %s", seg.id, exc)
                seg.status = "error"
                seg.validation_notes.append(f"Error de traducción: {exc}")

        context.setdefault("memory_log", []).append(
            f"[translator] {len(segments)} segmentos procesados ({source_lang} -> {target_lang})."
        )
        return context

    def retranslate_segment(self, seg: Segment, source_lang: str, target_lang: str, reason: str) -> None:
        """Reintenta un único segmento marcado como sospechoso por el validador."""
        try:
            seg.translated = self._translate_one(
                seg.original, source_lang, target_lang, retry_reason=reason
            )
        except Exception as exc:  # noqa: BLE001
            logger.error("[Translator] Fallo en reintento del segmento %d: %s", seg.id, exc)
            seg.status = "error"
            seg.validation_notes.append(f"Error en reintento: {exc}")


# --- AGENTE 3: VALIDADOR ---
_TECH_WHITELIST = {
    "ok", "internet", "email", "online", "web", "app", "gps", "usb", "pdf",
    "arduino", "ide", "led", "iot", "nano", "protoboard", "bootloader", "driver",
    "software", "hardware", "pin", "monitor", "python", "linux", "windows",
    "http", "https", "url", "bluetooth", "wifi", "microcontroller", "chip",
}


class ValidatorAgent:
    name = "validator"

    def _length_ratio_suspicious(self, original: str, translated: str) -> bool:
        len_o, len_t = len(original.strip()), len(translated.strip())
        if len_o < 30:  # títulos o líneas cortas varían naturalmente
            return False
        diff_ratio = abs(len_t - len_o) / len_o
        return diff_ratio > settings.LENGTH_DIFF_THRESHOLD

    def _find_untranslated_words(self, original: str, translated: str, source_lang: str, target_lang: str) -> List[str]:
        # Si el idioma origen y destino son iguales, las palabras idénticas son esperadas
        if source_lang.lower().strip() == target_lang.lower().strip():
            return []

        orig_clean = re.sub(r"https?://\S+|[^\w\s]", " ", original)
        trans_clean = re.sub(r"https?://\S+|[^\w\s]", " ", translated)

        orig_words = [w for w in orig_clean.split() if len(w) >= settings.MIN_UNTRANSLATED_WORD_LEN]
        if not orig_words:
            return []

        trans_words_set = {w.lower() for w in trans_clean.split()}
        
        suspicious = []
        for w in orig_words:
            wl = w.lower()
            if wl in _TECH_WHITELIST:
                continue
            # Nombres propios (empiezan con Mayúscula en el original) se ignoran
            if w[0].isupper():
                continue
            if wl in trans_words_set:
                suspicious.append(wl)

        unique_suspicious = sorted(set(suspicious))
        # Solo se considera sospechoso si una fracción significativa (>40%) de palabras no traducidas persiste
        if len(unique_suspicious) >= 3 and (len(suspicious) / len(orig_words)) > 0.40:
            return unique_suspicious
        return []

    def validate_segment(self, seg: Segment, source_lang: str = "auto", target_lang: str = "es") -> bool:
        seg.validation_notes = []
        problems = []

        if not seg.translated or not seg.translated.strip():
            problems.append("La traducción está vacía.")

        if self._length_ratio_suspicious(seg.original, seg.translated):
            len_o, len_t = len(seg.original.strip()), len(seg.translated.strip())
            ratio = abs(len_t - len_o) / max(len_o, 1)
            problems.append(
                f"Diferencia de longitud ({ratio:.0%}) supera el umbral configurado."
            )

        untranslated = self._find_untranslated_words(seg.original, seg.translated, source_lang, target_lang)
        if untranslated:
            problems.append(
                "Posibles palabras sin traducir: " + ", ".join(untranslated[:6])
            )

        if problems:
            seg.status = "sospechoso"
            seg.validation_notes = problems
            return False

        seg.status = "ok"
        return True

    def run(self, context: dict, translator: Optional[TranslatorAgent] = None) -> dict:
        segments: List[Segment] = context.get("segments", [])
        source_lang = context.get("source_lang", settings.DEFAULT_SOURCE_LANG)
        target_lang = context.get("target_lang", settings.DEFAULT_TARGET_LANG)
        translator = translator or context.get("translator_agent")

        error_count = 0
        for seg in segments:
            if seg.status == "error":
                error_count += 1
                continue

            passed = self.validate_segment(seg, source_lang, target_lang)
            while not passed and seg.retries < settings.MAX_RETRIES and translator is not None:
                seg.retries += 1
                reason = "; ".join(seg.validation_notes)
                logger.info(
                    "[Validator] Segmento %d sospechoso (intento %d/%d): %s",
                    seg.id, seg.retries, settings.MAX_RETRIES, reason,
                )
                translator.retranslate_segment(seg, source_lang, target_lang, reason)
                if seg.status == "error":
                    break
                passed = self.validate_segment(seg, source_lang, target_lang)

            if seg.status != "ok":
                error_count += 1

        total = len(segments) or 1
        error_rate = error_count / total
        context["error_rate"] = error_rate
        context.setdefault("memory_log", []).append(
            f"[validador] {total - error_count}/{total} segmentos OK (tasa de error {error_rate:.1%})."
        )
        logger.info("[Validator] Tasa de error del archivo: %.1f%%", error_rate * 100)
        return context


# --- AGENTE 4: ALINEADOR ---
class AlignerAgent:
    name = "aligner"

    def __init__(self, api_key: str = None):
        self._embeddings = None
        self.api_key = api_key or settings.GEMINI_API_KEY

    def _get_embeddings_client(self):
        if self._embeddings is None:
            from langchain_google_genai import GoogleGenerativeAIEmbeddings
            self._embeddings = GoogleGenerativeAIEmbeddings(
                model=settings.GEMINI_EMBEDDING_MODEL,
                google_api_key=self.api_key,
            )
        return self._embeddings

    def _align_by_position(self, segments: List[Segment]) -> None:
        for seg in segments:
            seg.color = color_for_index(seg.id)

    def _align_by_semantics(self, segments: List[Segment]) -> None:
        client = self._get_embeddings_client()
        originals = [s.original for s in segments]
        translations = [s.translated or s.original for s in segments]

        try:
            emb_orig = np.array(client.embed_documents(originals))
            emb_trans = np.array(client.embed_documents(translations))
        except Exception as exc:  # noqa: BLE001
            logger.warning(
                "[Aligner] Falló cálculo de embeddings (%s); usando alineación por posición.", exc
            )
            self._align_by_position(segments)
            return

        def cosine(a, b):
            na, nb = np.linalg.norm(a), np.linalg.norm(b)
            if na == 0 or nb == 0:
                return 0.0
            return float(np.dot(a, b) / (na * nb))

        for i, seg in enumerate(segments):
            sim = cosine(emb_orig[i], emb_trans[i])
            seg.color = color_for_index(i)
            if sim < settings.SEMANTIC_SIMILARITY_THRESHOLD:
                seg.validation_notes.append(
                    f"Similitud semántica baja ({sim:.2f}); revisar correspondencia."
                )
                if seg.status == "ok":
                    seg.status = "sospechoso"

    def run(self, context: dict) -> dict:
        segments: List[Segment] = context.get("segments", [])
        mode = context.get("alignment_mode", settings.ALIGNMENT_MODE)

        if mode == "semantic":
            self._align_by_semantics(segments)
        else:
            self._align_by_position(segments)

        context.setdefault("memory_log", []).append(
            f"[aligner] {len(segments)} segmentos alineados (modo='{mode}')."
        )
        return context


# ==============================================================================
# 7. ORQUESTADOR PRINCIPAL DEL PIPELINE
# ==============================================================================
class ConversationBufferMemory:
    """Buffer de memoria para registrar la trazabilidad y eventos del pipeline."""
    def __init__(self, return_messages: bool = True):
        self.return_messages = return_messages
        self.messages: list[str] = []

    def save_context(self, inputs: dict, outputs: dict):
        inp = inputs.get("input", "")
        out = outputs.get("output", "")
        self.messages.append(f"{inp}: {out}")

    def get_transcript(self) -> str:
        return "\n".join(self.messages)


ProgressCallback = Optional[Callable[[str, float], None]]


class TranslationOrchestrator:
    def __init__(self, source_lang: str = None, target_lang: str = None,
                 alignment_mode: str = None, api_key: str = None):
        self.source_lang = source_lang or settings.DEFAULT_SOURCE_LANG
        self.target_lang = target_lang or settings.DEFAULT_TARGET_LANG
        self.alignment_mode = alignment_mode or settings.ALIGNMENT_MODE
        self.api_key = api_key or settings.GEMINI_API_KEY

        # Memoria compartida del pipeline
        self.memory = ConversationBufferMemory(return_messages=True)

        self.extractor = ExtractorAgent()
        self.translator = TranslatorAgent(api_key=self.api_key)
        self.validator = ValidatorAgent()
        self.aligner = AlignerAgent(api_key=self.api_key)

    def _log_to_memory(self, context: dict, filename: str):
        for line in context.get("memory_log", []):
            self.memory.save_context(
                {"input": f"[{filename}] evento"},
                {"output": line},
            )

    def process_file(self, filename: str, file_bytes: bytes,
                     on_progress: ProgressCallback = None) -> dict:
        context = {
            "filename": filename,
            "file_bytes": file_bytes,
            "source_lang": self.source_lang,
            "target_lang": self.target_lang,
            "alignment_mode": self.alignment_mode,
            "translator_agent": self.translator,
            "memory_log": [],
            "started_at": time.time(),
        }

        stages = [
            ("Extrayendo texto", lambda ctx: self.extractor.run(ctx), 0.10, 0.25),
            ("Traduciendo con Gemini", lambda ctx: self.translator.run(ctx, on_progress=on_progress), 0.25, 0.70),
            ("Validando traducción", lambda ctx: self.validator.run(ctx, self.translator), 0.70, 0.85),
            ("Alineando segmentos", lambda ctx: self.aligner.run(ctx), 0.85, 1.0),
        ]

        for stage_name, fn, start_frac, end_frac in stages:
            if on_progress:
                on_progress(stage_name, start_frac)
            if context.get("error") and stage_name != "Extrayendo texto":
                break
            try:
                context = fn(context)
            except Exception as exc:  # noqa: BLE001
                logger.exception("Fallo inesperado en etapa '%s' para %s", stage_name, filename)
                context["error"] = f"Error en etapa '{stage_name}': {exc}"
                break
            if on_progress:
                on_progress(f"{stage_name} completado", end_frac)

        context["finished_at"] = time.time()
        context["duration_seconds"] = context["finished_at"] - context["started_at"]
        self._log_to_memory(context, filename)

        if context.get("error"):
            context["file_status"] = "error"
        elif context.get("error_rate", 0) > 0.10:
            context["file_status"] = "traducido_con_advertencias"
        else:
            context["file_status"] = "traducido"

        return context

    def get_memory_transcript(self) -> str:
        """Devuelve el historial acumulado de todos los agentes (para auditoría)."""
        return self.memory.get_transcript()


# ==============================================================================
# 8. INTERFAZ STREAMLIT (UI)
# ==============================================================================
st.set_page_config(
    page_title="Traducción por Lotes con Agentes (Gemini)",
    page_icon="🌐",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Inyección de estilos CSS avanzados (Glassmorphism + Modern Typography)
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Outfit:wght@400;600;700;800&family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');

    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }

    h1, h2, h3, h4, h5, h6 {
        font-family: 'Outfit', sans-serif !important;
        letter-spacing: -0.02em;
    }

    /* Hero Glassmorphic Card */
    .hero-banner {
        background: linear-gradient(135deg, rgba(30, 41, 59, 0.7) 0%, rgba(15, 23, 42, 0.85) 100%);
        border: 1px solid rgba(99, 102, 241, 0.25);
        border-radius: 16px;
        padding: 24px 28px;
        margin-bottom: 24px;
        box-shadow: 0 10px 30px 0 rgba(0, 0, 0, 0.35);
        backdrop-filter: blur(12px);
    }

    .hero-title {
        font-size: 2.1rem;
        font-weight: 800;
        background: linear-gradient(135deg, #818CF8 0%, #C084FC 50%, #F472B6 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 8px;
    }

    .agent-pipeline {
        display: flex;
        flex-wrap: wrap;
        gap: 10px;
        margin-top: 14px;
    }

    .agent-pill {
        display: inline-flex;
        align-items: center;
        gap: 6px;
        padding: 6px 14px;
        background: rgba(255, 255, 255, 0.05);
        border: 1px solid rgba(255, 255, 255, 0.12);
        border-radius: 30px;
        font-size: 0.82rem;
        font-weight: 600;
        color: #E2E8F0;
        transition: all 0.2s ease;
    }

    .agent-pill:hover {
        background: rgba(99, 102, 241, 0.25);
        border-color: rgba(99, 102, 241, 0.5);
        transform: translateY(-2px);
        box-shadow: 0 4px 12px rgba(99, 102, 241, 0.3);
    }

    /* Metric cards styling */
    [data-testid="stMetric"] {
        background: linear-gradient(135deg, rgba(30, 41, 59, 0.5), rgba(15, 23, 42, 0.7));
        border: 1px solid rgba(255, 255, 255, 0.08);
        border-radius: 12px;
        padding: 12px 16px;
        box-shadow: 0 4px 14px rgba(0,0,0,0.2);
    }

    /* Custom scrollbars */
    ::-webkit-scrollbar {
        width: 7px;
        height: 7px;
    }
    ::-webkit-scrollbar-track {
        background: rgba(15, 23, 42, 0.4);
        border-radius: 10px;
    }
    ::-webkit-scrollbar-thumb {
        background: rgba(99, 102, 241, 0.45);
        border-radius: 10px;
    }
    ::-webkit-scrollbar-thumb:hover {
        background: rgba(99, 102, 241, 0.75);
    }

    .segment-pill:hover {
        transform: translateY(-2px);
        box-shadow: 0 6px 16px rgba(0,0,0,0.22) !important;
    }
</style>
""", unsafe_allow_html=True)

STATUS_LABELS = {
    "pendiente": "⏳ Pendiente",
    "procesando": "🔄 Procesando",
    "traducido": "✅ Traducido",
    "traducido_con_advertencias": "⚠️ Traducido (con advertencias)",
    "error": "❌ Error",
}

LANGUAGE_OPTIONS_SOURCE = {
    "auto": "🌐 Detección automática (auto)",
    "es": "🇪🇸 Español (es)",
    "en": "🇬🇧 Inglés (en)",
    "fr": "🇫🇷 Francés (fr)",
    "de": "🇩🇪 Alemán (de)",
    "it": "🇮🇹 Italiano (it)",
    "pt": "🇵🇹 Portugués (pt)",
    "zh": "🇨🇳 Chino (zh)",
    "ja": "🇯🇵 Japonés (ja)",
    "ru": "🇷🇺 Ruso (ru)",
    "ar": "🇸🇦 Árabe (ar)",
}

LANGUAGE_OPTIONS_TARGET = {
    "es": "🇪🇸 Español (es)",
    "en": "🇬🇧 Inglés (en)",
    "fr": "🇫🇷 Francés (fr)",
    "de": "🇩🇪 Alemán (de)",
    "it": "🇮🇹 Italiano (it)",
    "pt": "🇵🇹 Portugués (pt)",
    "zh": "🇨🇳 Chino (zh)",
    "ja": "🇯🇵 Japonés (ja)",
    "ru": "🇷🇺 Ruso (ru)",
    "ar": "🇸🇦 Árabe (ar)",
}


def create_batch_zip(results: dict, export_format: str) -> bytes:
    """Empaqueta todos los archivos procesados con éxito en un único .zip."""
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for fname, context in results.items():
            segments = context.get("segments", [])
            if segments and not context.get("error"):
                # Sincronizar con posibles ediciones manuales en session_state
                for seg in segments:
                    val = st.session_state.get(f"edit_{fname}_{seg.id}", st.session_state.edited_texts.get((fname, seg.id), seg.translated))
                    seg.translated = val
                file_bytes_out = export_segments(segments, export_format)
                out_name = fname.rsplit(".", 1)[0] + f".{export_format}"
                zip_file.writestr(out_name, file_bytes_out)
    return zip_buffer.getvalue()


def init_session_state():
    st.session_state.setdefault("files_status", {})   # filename -> status str
    st.session_state.setdefault("results", {})        # filename -> context dict
    st.session_state.setdefault("uploaded_map", {})   # filename -> bytes
    st.session_state.setdefault("processing", False)
    st.session_state.setdefault("edited_texts", {})   # (filename, seg_id) -> str


init_session_state()

# --- Barra lateral (Configuración) ---
with st.sidebar:
    st.markdown("### ⚙️ Configuración del Sistema")

    # Gestión de GEMINI_API_KEY
    env_api_key = settings.GEMINI_API_KEY.strip()
    if env_api_key:
        st.success("🔒 Clave API cargada del entorno (.env)")
        with st.expander("🔑 Cambiar API Key para esta sesión"):
            override_key = st.text_input(
                "Nueva GEMINI_API_KEY (opcional)",
                value="",
                type="password",
                help="Deja este campo en blanco para usar la clave cargada del .env.",
            )
        api_key_input = override_key.strip() if override_key.strip() else env_api_key
    else:
        api_key_input = st.text_input(
            "GEMINI_API_KEY (Obligatorio)",
            value="",
            type="password",
            help="Introduce tu API Key de Google Gemini.",
        )
        if not api_key_input:
            st.warning("⚠️ Debes ingresar tu GEMINI_API_KEY para poder procesar.")

    st.markdown("---")
    st.markdown("#### 🌐 Parámetros de Traducción")

    col_a, col_b = st.columns(2)
    source_keys = list(LANGUAGE_OPTIONS_SOURCE.keys())
    target_keys = list(LANGUAGE_OPTIONS_TARGET.keys())

    default_src_idx = source_keys.index(settings.DEFAULT_SOURCE_LANG) if settings.DEFAULT_SOURCE_LANG in source_keys else 0
    default_tgt_idx = target_keys.index(settings.DEFAULT_TARGET_LANG) if settings.DEFAULT_TARGET_LANG in target_keys else 0

    with col_a:
        source_lang = st.selectbox(
            "Idioma origen",
            options=source_keys,
            index=default_src_idx,
            format_func=lambda k: LANGUAGE_OPTIONS_SOURCE.get(k, k),
            help="Selecciona el idioma del documento o 'auto' para detección automática.",
        )
    with col_b:
        target_lang = st.selectbox(
            "Idioma destino",
            options=target_keys,
            index=default_tgt_idx,
            format_func=lambda k: LANGUAGE_OPTIONS_TARGET.get(k, k),
            help="Selecciona el idioma al que se traducirá el documento.",
        )

    alignment_mode = st.selectbox(
        "Modo de alineación",
        options=["position", "semantic"],
        index=0 if settings.ALIGNMENT_MODE == "position" else 1,
        help="'position': por orden de párrafo (rápido). "
             "'semantic': valida además con embeddings de Gemini.",
    )

    export_format = st.selectbox("Formato de exportación", options=["txt", "docx", "pdf"], index=1)

    st.markdown("---")
    st.caption(
        f"**Umbral de longitud:** ±{settings.LENGTH_DIFF_THRESHOLD:.0%} · "
        f"**Reintentos:** {settings.MAX_RETRIES} · "
        f"**Rate limit:** {settings.MAX_CALLS_PER_MINUTE} req/min"
    )

# --- Cabecera Principal (Hero Banner) ---
st.markdown("""
<div class="hero-banner">
    <div class="hero-title">🌐 Sistema de Traducción por Lotes con Agentes</div>
    <div style="color: #94A3B8; font-size: 0.95rem; line-height: 1.5;">
        Sube múltiples documentos (<code>.txt</code>, <code>.docx</code>, <code>.pdf</code>) y un equipo coordinado de 
        <strong>4 agentes inteligentes</strong> procesará el lote con validación cruzada y verificación visual por colores en tiempo real.
    </div>
    <div class="agent-pipeline">
        <span class="agent-pill">📄 1. Extractor</span>
        <span class="agent-pill">🌐 2. Traductor (Gemini Flash-Lite)</span>
        <span class="agent-pill">🛡️ 3. Validador Automático</span>
        <span class="agent-pill">🔗 4. Alineador Semántico</span>
    </div>
</div>
""", unsafe_allow_html=True)

# --- 1. Carga de Archivos ---
uploaded_files = st.file_uploader(
    "📁 Arrastra o selecciona tus archivos para procesar en lote:",
    type=["txt", "docx", "pdf"],
    accept_multiple_files=True,
)

if uploaded_files:
    for f in uploaded_files:
        if f.name not in st.session_state.uploaded_map:
            st.session_state.uploaded_map[f.name] = f.getvalue()
            st.session_state.files_status.setdefault(f.name, "pendiente")

# Tabla de estado de archivos
if st.session_state.uploaded_map:
    st.subheader("📋 Estado de Archivos del Lote")
    status_rows = [
        {"Archivo": fname, "Tamaño": f"{len(st.session_state.uploaded_map[fname])/1024:.1f} KB", "Estado": STATUS_LABELS.get(status, status)}
        for fname, status in st.session_state.files_status.items()
    ]
    st.table(status_rows)

# --- 2. Procesamiento del Lote ---
src_label = LANGUAGE_OPTIONS_SOURCE.get(source_lang, source_lang)
tgt_label = LANGUAGE_OPTIONS_TARGET.get(target_lang, target_lang)
st.info(f"🌐 **Ruta activa:** {src_label} ➔ **{tgt_label}** *(Ajusta los idiomas en la barra lateral)*")

run_col1, run_col2, run_col3 = st.columns([2, 2, 2])
with run_col1:
    start_clicked = st.button(
        "🚀 Procesar Lote",
        disabled=st.session_state.processing or not st.session_state.uploaded_map,
        use_container_width=True,
        type="primary",
    )
with run_col2:
    if st.button("🗑️ Limpiar Todo / Nuevo Lote", disabled=st.session_state.processing, use_container_width=True):
        st.session_state.files_status.clear()
        st.session_state.results.clear()
        st.session_state.uploaded_map.clear()
        st.session_state.edited_texts.clear()
        for k in list(st.session_state.keys()):
            if str(k).startswith("edit_") or str(k).startswith("download_"):
                del st.session_state[k]
        st.rerun()

if start_clicked:
    if not api_key_input:
        st.error("⚠️ Debes configurar tu GEMINI_API_KEY en la barra lateral antes de procesar.")
    else:
        st.session_state.processing = True

        orchestrator = TranslationOrchestrator(
            source_lang=source_lang,
            target_lang=target_lang,
            alignment_mode=alignment_mode,
            api_key=api_key_input,
        )

        files_to_process = list(st.session_state.uploaded_map.keys())

        overall_progress = st.progress(0.0, text="Iniciando procesamiento del lote...")
        live_area = st.container()

        total = len(files_to_process) or 1
        for idx, fname in enumerate(files_to_process):
            st.session_state.files_status[fname] = "procesando"

            # Limpiar cualquier edición previa o llaves cacheadas de este archivo al reprocesar
            for key in list(st.session_state.edited_texts.keys()):
                if key[0] == fname:
                    del st.session_state.edited_texts[key]
            for s_key in list(st.session_state.keys()):
                if str(s_key).startswith(f"edit_{fname}_"):
                    del st.session_state[s_key]

            def _on_progress(stage_name: str, frac: float, _fname=fname, _idx=idx):
                overall_progress.progress(
                    min((_idx + frac) / total, 1.0),
                    text=f"[{_fname}] {stage_name}...",
                )

            try:
                file_bytes = st.session_state.uploaded_map[fname]
                context = orchestrator.process_file(fname, file_bytes, on_progress=_on_progress)
            except Exception as exc:  # noqa: BLE001
                logger.error("Fallo procesando %s: %s\n%s", fname, exc, traceback.format_exc())
                context = {"error": str(exc), "file_status": "error", "segments": []}

            st.session_state.results[fname] = context
            st.session_state.files_status[fname] = context.get("file_status", "error")

            with live_area:
                if context.get("file_status") != "error":
                    st.success(f"✅ Completado: {fname}")
                else:
                    st.error(f"❌ Error en {fname}: {context.get('error')}")

        overall_progress.progress(1.0, text="Lote completado con éxito.")
        st.session_state.processing = False
        st.rerun()

# --- 3. Revisión, Verificación Cruzada y Edición ---
if st.session_state.results:
    st.markdown("### 🔍 Resultados y Verificación Cruzada")

    # Botón global de descarga de lote en ZIP si hay archivos procesados con éxito
    valid_results = [
        c for c in st.session_state.results.values()
        if c.get("segments") and not c.get("error")
    ]
    if valid_results:
        zip_bytes = create_batch_zip(st.session_state.results, export_format)
        st.download_button(
            label=f"📦 Descargar Lote Completo (.ZIP con archivos .{export_format})",
            data=zip_bytes,
            file_name=f"traducciones_lote_{export_format}.zip",
            mime="application/zip",
            key="download_all_zip",
            type="primary",
        )
        st.markdown("")

    for fname, context in st.session_state.results.items():
        status = context.get("file_status", "pendiente")
        with st.expander(f"{STATUS_LABELS.get(status, status)} — 📄 {fname}", expanded=(status != "error")):

            if context.get("error"):
                st.error(context["error"])
                continue

            segments = context.get("segments", [])
            if not segments:
                st.warning("No se extrajeron segmentos de este archivo.")
                continue

            error_rate = context.get("error_rate", 0.0)
            duration = context.get("duration_seconds", 0.0)

            # Métricas visuales del archivo (en lugar de solo un caption plano)
            m1, m2, m3, m4 = st.columns(4)
            with m1:
                st.metric(label="📊 Segmentos Extraídos", value=len(segments))
            with m2:
                st.metric(label="🛡️ Tasa de Advertencia", value=f"{error_rate:.1%}")
            with m3:
                st.metric(label="⚡ Tiempo Total", value=f"{duration:.1f}s")
            with m4:
                st.metric(label="🎯 Estado Final", value=STATUS_LABELS.get(status, status))

            st.markdown("---")

            # Vista previa sincronizada de solo lectura con scroll interno
            st.markdown("##### 📖 Vista Previa Sincronizada (Verificación Cruzada por Color)")
            originals_text = [s.original for s in segments]
            translated_text_display = [
                st.session_state.get(f"edit_{fname}_{s.id}", st.session_state.edited_texts.get((fname, s.id), s.translated or s.original))
                for s in segments
            ]
            colors = [s.color for s in segments]

            col_orig, col_trans = st.columns(2)
            with col_orig:
                st.markdown(f"**📄 Documento Original ({source_lang.upper()})**")
                with st.container(height=390):
                    st.markdown(render_highlighted_html(originals_text, colors), unsafe_allow_html=True)

            with col_trans:
                st.markdown(f"**🌐 Traducción Generada ({target_lang.upper()})**")
                with st.container(height=390):
                    st.markdown(render_highlighted_html(translated_text_display, colors), unsafe_allow_html=True)

            # Edición manual por segmento en contenedor colapsable separado
            with st.expander("✏️ Editor Manual de Segmentos (Opcional)", expanded=False):
                st.caption("Puedes modificar el texto de cualquier segmento antes de exportar. Los cambios se actualizarán en la descarga.")
                with st.container(height=320):
                    for seg in segments:
                        key = f"edit_{fname}_{seg.id}"
                        default_value = st.session_state.get(
                            key, st.session_state.edited_texts.get((fname, seg.id), seg.translated or seg.original)
                        )
                        label_warning = f"({'⚠️ ' + '; '.join(seg.validation_notes) if seg.validation_notes else 'OK'})"
                        new_value = st.text_area(
                            label=f"Segmento {seg.id + 1} {label_warning}",
                            value=default_value,
                            key=key,
                            height=90,
                        )
                        st.session_state.edited_texts[(fname, seg.id)] = new_value
                        seg.translated = new_value

            # --- Exportación individual ---
            try:
                # Sincronizar los segmentos con las ediciones
                for seg in segments:
                    val = st.session_state.get(f"edit_{fname}_{seg.id}", st.session_state.edited_texts.get((fname, seg.id), seg.translated))
                    seg.translated = val
                file_bytes_out = export_segments(segments, export_format)
                mime_map = {
                    "txt": "text/plain",
                    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "pdf": "application/pdf",
                }
                out_name = fname.rsplit(".", 1)[0] + f".{export_format}"
                st.download_button(
                    label=f"⬇️ Descargar '{out_name}' ({export_format.upper()})",
                    data=file_bytes_out,
                    file_name=out_name,
                    mime=mime_map[export_format],
                    key=f"download_{fname}",
                )
            except UnsupportedFormatError as exc:
                st.error(str(exc))

    with st.expander("🧠 Registro de Auditoría de Agentes (Memoria Compartida)"):
        st.caption(
            "Historial de eventos, etapas y decisiones generadas por cada agente en la memoria del pipeline."
        )
        for fname, context in st.session_state.results.items():
            for line in context.get("memory_log", []):
                st.text(f"[{fname}] {line}")
